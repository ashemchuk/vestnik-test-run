from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw

root = Path(__file__).resolve().parents[1] / 'test-data'
root.mkdir(exist_ok=True)
image = Image.new('RGB', (900, 300), 'white')
draw = ImageDraw.Draw(image)
draw.line([(70, 30), (70, 260), (850, 260)], fill='black', width=3)
points = [(70 + n * 9, 260 - count * .065) for n, count in [(10, 45), (20, 190), (40, 780), (80, 3160)]]
draw.line(points, fill='#333333', width=4)
for (x, y), n in zip(points, [10, 20, 40, 80]):
    draw.ellipse((x-5, y-5, x+5, y+5), fill='black')
    draw.text((x-5, 275), str(n), fill='black')
draw.text((15, 10), 'L', fill='black')
draw.text((860, 255), 'n', fill='black')
draw.text((20, 45), '3160', fill='black')
image.save(root / 'comparisons.png')

doc = Document()
section = doc.sections[0]
section.page_height, section.page_width = Cm(29.7), Cm(21)
section.top_margin = section.bottom_margin = Cm(2)
section.left_margin = section.right_margin = Cm(2)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1
doc.styles['Title'].font.name = 'Times New Roman'
doc.styles['Title'].font.size = Pt(16)
doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
doc.add_paragraph('Сравнение алгоритмов поиска дубликатов', 'Title')
doc.add_paragraph('Аннотация. Рассматриваются попарное сравнение элементов и поиск повторов с помощью множества. Для небольших массивов оба подхода дают одинаковый результат, но число операций различается. Цель работы — показать зависимость числа сравнений от размера массива.')
doc.add_paragraph('Ключевые слова: алгоритмы, дубликаты, сложность вычислений.')
doc.add_paragraph('Метод. Попарный алгоритм сравнивает каждый элемент со всеми последующими. Для массива без повторов число сравнений составляет L = n(n − 1) / 2. Алгоритм с множеством последовательно проверяет наличие элемента и добавляет его в множество; при стандартных предположениях о хешировании ожидаемое время линейно.')
doc.add_paragraph('Результаты. Для массивов из 10, 20, 40 и 80 различных элементов попарный алгоритм выполняет соответственно 45, 190, 780 и 3160 сравнений. Эти значения получены по формуле и не являются измерениями времени работы.')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(root / 'comparisons.png'), width=Cm(13))
doc.add_paragraph('Рисунок 1. Число сравнений в массиве без повторов.')
doc.add_paragraph('Заключение. Попарное сравнение требует квадратичного числа операций. Для больших массивов применение множества сокращает ожидаемую трудоёмкость. Для воспроизводимого измерения времени необходимо дополнительно зафиксировать реализацию и среду исполнения.')
for key in ['author', 'last_modified_by', 'subject', 'keywords', 'comments']:
    setattr(doc.core_properties, key, '')
doc.core_properties.title = 'Сравнение алгоритмов поиска дубликатов'
doc.save(root / 'article-a-v1.docx')
print(root / 'article-a-v1.docx')
