from pathlib import Path
import json,secrets
from docx import Document
from docx.shared import Pt,Cm
root=Path('/Users/shemchuk/Projects/vestnik'); data=root/'test-data/full-run'; data.mkdir(parents=True,exist_ok=True)
accounts=[]
for username,given,family,roles in [('chiefeditor','Главный','Редактор',['Journal editor']),('reviewer3','Третий','Рецензент',['Reviewer']),('reviewer4','Четвёртый','Рецензент',['Reviewer']),('reviewer5','Пятый','Рецензент',['Reviewer','Editorial Board Member']),('author2','Второй','Автор',['Author']),('author3','Третий','Автор',['Author'])]:
 accounts.append(dict(username=username,givenName=given,familyName=family,roles=roles,role=roles[0],email=username+'@example.test',password=secrets.token_urlsafe(20)))
p=root/'.local-full-users.json'; p.write_text(json.dumps(accounts,ensure_ascii=False,indent=2)); p.chmod(0o600)
(root/'.gitignore').write_text((root/'.gitignore').read_text()+'\n.local-full-users.json\nbackups/\n')
(root/'backups/before-full-run.sql').chmod(0o600)
for fname,version in [('A-original',0),('A-formatted',1),('A-v2',2),('C-styled',3)]:
 d=Document(); d.styles['Normal'].font.name='Times New Roman'; d.styles['Normal'].font.size=Pt(11 if version else 14)
 d.styles['Normal'].paragraph_format.line_spacing=1
 s=d.sections[0]; s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Cm(2)
 title='Проверка поиска дубликатов в массивах' if version<3 else 'Короткая заметка о двоичном поиске'
 p=d.add_paragraph(title, 'Title' if version==3 else 'Normal')
 if version<3: p.runs[0].bold=True
 d.add_paragraph('Тестовая рукопись для проверки редакционного процесса OJS. Сведения об авторах указаны отдельно в метаданных системы.')
 d.add_paragraph('Метод' if version==3 else 'Метод и результаты', 'Heading 1' if version==3 else 'Normal')
 d.add_paragraph('Попарный поиск требует L = n(n − 1) / 2 сравнений. Множество обеспечивает ожидаемое линейное время при стандартных предположениях о хешировании.' if version<3 else 'Двоичный поиск уменьшает диапазон пополам на каждой итерации. Для отсортированного массива требуется O(log n) сравнений.')
 if version<3:
  d.add_picture(str(root/'test-data/comparisons-v2.png'),width=Cm(13)); d.add_paragraph('Рисунок 1. Число сравнений в зависимости от длины массива n.')
 if version==2:
  d.add_paragraph('Исправления по рецензиям'); d.add_paragraph('Добавлены граничные случаи: пустой массив, один элемент, повторяющиеся значения. Псевдокод: seen = set(); для x: если x in seen, вернуть повтор; иначе добавить x. Ожидаемая память O(n).')
 for prop in ['author','last_modified_by','subject','keywords','comments']: setattr(d.core_properties,prop,'')
 d.core_properties.title=title; d.save(data/(fname+'.docx'))
(data/'A-response-v2.txt').write_text('Ответ на рецензии\n1. Уточнены граничные случаи.\n2. Добавлены псевдокод и оценка памяти.\n3. Формула и график проверены.\n',encoding='utf-8')
(data/'B.tex').write_text(r'''\documentclass{article}
\usepackage{graphicx}
\begin{document}
\title{Queue scheduling experiment}\maketitle
This anonymous test manuscript describes FIFO queue scheduling. The mean waiting time is $W=\sum w_i/n$.
\includegraphics[width=0.7\textwidth]{figure.png}
\bibliographystyle{plain}\bibliography{B}
\end{document}
''')
(data/'B.bib').write_text('@book{test, title={Algorithms}, author={Test Reference}, year={2026}}\n')
import shutil,zipfile
shutil.copy(root/'test-data/comparisons-v2.png',data/'figure.png')
with zipfile.ZipFile(data/'B-latex.zip','w') as z:
 for f in ['B.tex','B.bib','figure.png']: z.write(data/f,f)
(data/'blocked-upload.php').write_text('<?php /* inert test file; contains no executable actions */ ?>')
(data/'blocked-upload.html').write_text('<!doctype html><title>Inert upload test</title><p>test</p>')
print('Prepared 6 accounts, four DOCX fixtures and LaTeX archive')
