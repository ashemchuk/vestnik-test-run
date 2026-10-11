from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont

root = Path(__file__).resolve().parents[1] / 'test-data'
doc = Document(root / 'article-a-v1.docx')
image = Image.new('RGB', (1000, 360), 'white')
draw = ImageDraw.Draw(image)
font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 20)
draw.line([(120, 70), (120, 285), (930, 285)], fill='black', width=3)
draw.text((20, 12), 'Число сравнений L', fill='black', font=font)
draw.text((540, 325), 'Число элементов массива n', fill='black', font=font)
points = [(120 + n * 9.5, 285 - count * .06) for n, count in [(10,45),(20,190),(40,780),(80,3160)]]
draw.line(points, fill='black', width=3)
for (x,y), n, count in zip(points,[10,20,40,80],[45,190,780,3160]):
    draw.ellipse((x-5,y-5,x+5,y+5),fill='black')
    draw.text((x-12,292),str(n),fill='black',font=font)
    draw.text((x-20,y-27),str(count),fill='black',font=font)
image.save(root / 'comparisons-v2.png')
for p in doc.paragraphs:
    if p._p.xpath('.//w:drawing'):
        p.clear()
        p.add_run().add_picture(str(root / 'comparisons-v2.png'), width=Cm(13))
    if p.text.startswith('Рисунок 1.'):
        p.text = 'Рисунок 1. Расчётное число сравнений L в массиве из n различных элементов. Время работы не измерялось.'
doc.add_page_break()
doc.add_paragraph('Описание алгоритмов и пример', 'Title')
doc.add_paragraph('Оба алгоритма отвечают на вопрос, есть ли в массиве хотя бы два равных элемента. Они возвращают логическое значение и завершаются сразу после обнаружения повтора.')
doc.add_paragraph('Попарный алгоритм')
for line in ['для i от 0 до n − 2:', '    для j от i + 1 до n − 1:', '        если A[i] = A[j]: вернуть Истина', 'вернуть Ложь']:
    p=doc.add_paragraph(line)
    p.paragraph_format.space_after=Pt(0)
    for run in p.runs: run.font.name='Courier New'; run.font.size=Pt(10)
doc.add_paragraph('Алгоритм с множеством')
for line in ['S := пустое множество', 'для каждого x в A:', '    если x содержится в S: вернуть Истина', '    добавить x в S', 'вернуть Ложь']:
    p=doc.add_paragraph(line)
    p.paragraph_format.space_after=Pt(0)
    for run in p.runs: run.font.name='Courier New'; run.font.size=Pt(10)
doc.add_paragraph('Для массива [3, 1, 3, 2] оба алгоритма возвращают Истина. Попарный алгоритм обнаруживает совпадение элементов с индексами 0 и 2; алгоритм с множеством обнаруживает повтор при обработке третьего элемента. Для [3, 1, 2] оба возвращают Ложь.')
doc.add_paragraph('При ожидаемой постоянной стоимости поиска и добавления в хеш-множество алгоритм выполняется за ожидаемое O(n) времени и требует O(n) дополнительной памяти. Это оценка при предположениях о хешировании, а не гарантия для любого набора данных. Попарный алгоритм в худшем случае выполняет n(n − 1) / 2 сравнений и использует O(1) дополнительной памяти.')
for key in ['author','last_modified_by','subject','keywords','comments']: setattr(doc.core_properties,key,'')
doc.save(root / 'article-a-v2.docx')
(root / 'response-to-reviewers-v2.txt').write_text('Ответ на замечания по статье №1\n\nРецензент 1\n1. Добавлены псевдокоды двух алгоритмов на странице 2.\n2. Уточнены предположения о хешировании, ожидаемая сложность O(n) и память O(n).\n\nРецензент 2\n1. Оси рисунка подписаны словами; показано расчётное число сравнений.\n2. Подпись явно сообщает, что время работы не измерялось.\n3. Добавлены массивы [3, 1, 3, 2] и [3, 1, 2] с результатами Истина и Ложь для обоих алгоритмов.\n\nПросим направить исправленные DOCX и PDF на повторную проверку.\n',encoding='utf-8')
print(root / 'article-a-v2.docx')
